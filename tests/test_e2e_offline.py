"""End-to-end: docx -> extract -> clean -> chunk -> index -> originality / gaps / chat (offline fakes)."""
import json
from pathlib import Path

from docx import Document

from src.chunk_for_rag import chunk_all
from src.clean import clean_all
from src.duplicate_detector import DuplicateDetector
from src.extract import extract_all
from src.gap_analyzer import GapAnalyzer
from src.rag_chat import RAGChatEngine
from src.vector_store import ProjectIQVectorStore

BODY = (
    "ABSTRACT\nWe detect plant leaf disease using a CNN with OpenCV on a Raspberry Pi camera and reach 94% accuracy. "
    + "The model classifies leaf images into healthy and diseased classes using transfer learning on a labelled dataset. " * 3
    + "\n\nMETHODOLOGY\nImages are captured, resized, normalised and passed to a ResNet classifier running on the device. "
    + "Training uses augmentation, dropout and early stopping to limit overfitting on the small dataset. " * 3
    + "\n\nFUTURE SCOPE\nA limitation is the small dataset and real-time latency on the edge device; further research could "
    + "extend the system to more crops and reduce computational cost with quantisation. " * 3
)


def _make_raw(root: Path):
    d = root / "DS_1" / "Group 4" / "Blackbook"
    d.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("A PROJECT REPORT ON")
    doc.add_paragraph("Plant Leaf Disease Detection")
    for para in BODY.split("\n\n"):
        doc.add_paragraph(para)
    doc.save(d / "final_report.docx")


def test_full_flow(tmp_path, offline, monkeypatch):
    from src.config import config
    monkeypatch.setattr(config, "min_relevance", 0.0)  # toy hash-embeddings score lower than MiniLM
    raw, ext, cln = tmp_path / "raw", tmp_path / "ext", tmp_path / "cln"
    _make_raw(raw)

    manifest = extract_all(raw, ext)
    assert manifest["extracted_files"] == 1
    meta = manifest["documents"][0]
    assert (meta["department"], meta["group_id"], meta["category"]) == ("DS_1", "Group 4", "blackbook")

    assert clean_all(ext, cln)["cleaned_documents"] == 1
    cleaned = json.loads(next(cln.glob("*_cleaned.json")).read_text(encoding="utf-8"))
    assert "authors" not in cleaned["document"].get("metadata", {})

    chunks_file = tmp_path / "chunks.jsonl"
    chunks = chunk_all(cln, chunks_file)
    assert chunks and all(c["token_count"] <= 520 for c in chunks)

    store = ProjectIQVectorStore(db_path=tmp_path / "db")
    assert store.index_chunks(chunks_file, prune=True) > 0
    assert store.index_chunks(chunks_file, prune=True) > 0  # idempotent: ids are stable
    assert store.count() == len([c for c in chunks if c["is_canonical"]])

    hits = store.query_similar("plant leaf disease CNN raspberry pi", n_results=3)
    assert hits and 0.0 < hits[0]["similarity"] <= 1.0
    assert hits[0]["similarity"] >= hits[-1]["similarity"]
    assert store.query_similar("x", department="NOPE") == []

    res = DuplicateDetector(store).check_originality("Plant leaf disease detection with CNN on Raspberry Pi")
    assert res["matching_projects"] and res["max_similarity_pct"] > 0
    assert DuplicateDetector(store).check_originality("")["verdict"] == "Empty Input"

    gaps = GapAnalyzer(store, api_key=None)
    gaps.api_key = None
    props = gaps.generate_novel_projects(domain_filter="plant disease", num_proposals=2)
    assert len(props) == 2 and all(3 <= p["feasibility_score"] <= 10 for p in props)

    chat = RAGChatEngine(store, api_key=None)
    chat._gemini_client = None
    out = chat.ask("What CNN methodology was used for leaf disease?")
    assert out["sources"] and out["engine_used"].startswith("Local")
