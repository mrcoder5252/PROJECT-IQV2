"""
End-to-End Integration Verification Test for Project-IQ (Tiers 3 & 4)
Tests full multi-format ingestion, cleaning, section chunking, vector indexing,
gap analysis ideation, and team matchmaking against a self-contained synthetic academic corpus.
"""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

from run_pipeline import run_full_pipeline
from src.gap_analyzer import GapAnalyzer
from src.team_matcher import TeamMatcher
from src.vector_store import ProjectIQVectorStore


def create_synthetic_college_corpus(raw_dir: Path):
    """
    Dynamically generates a realistic multi-format college project corpus
    (PDF, DOCX, PPTX) adhering to the standard college hierarchy
    (Department / Group / Document Type).
    """
    # 1. CS_1 / Group 1 / Blackbook / Smart_Agri_Blackbook.pdf
    grp1_bb_dir = raw_dir / "CS_1" / "Group 1" / "Blackbook"
    grp1_bb_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = grp1_bb_dir / "Smart_Agri_Blackbook.pdf"

    doc = fitz.open()
    # Page 1: Admin front-matter with student PII
    p1 = doc.new_page()
    p1_text = """SAVITRIBAI PHULE PUNE UNIVERSITY
DEPARTMENT OF COMPUTER ENGINEERING 2023-24
A PROJECT REPORT ON
SMART AGRICULTURE LEAF DISEASE DETECTION USING DEEP LEARNING

SUBMITTED BY:
VAIBHAV KADAM (Roll No: BE-COMP-42, PRN: 72018492K)
AMIT DESHMUKH (Roll No: BE-COMP-15, PRN: 72018431M)
Contact Email: vaibhav.kadam@college.edu | Phone: +91 9876543210
GUIDED BY:
PROF. S. K. SHARMA

THIS IS TO CERTIFY THAT the project report entitled "Smart Agriculture Leaf Disease Detection"
is a bonafide work carried out by the students in partial fulfillment of the degree.
"""
    p1.insert_text((50, 60), p1_text, fontsize=11)

    # Page 2: Abstract & Literature Survey
    p2 = doc.new_page()
    p2_text = """ABSTRACT
Agriculture productivity is heavily impacted by plant crop diseases. Early detection of leaf anomalies
can prevent massive harvest losses. In this project, we develop an automated deep convolutional neural
network framework for diagnosing tomato and potato leaf blights.

LITERATURE SURVEY
Previous studies by LeCun et al. [1] established the efficacy of deep convolution backpropagation.
Recent work by Vaswani et al. [3] demonstrated self-attention mechanisms in sequence modeling.
However, existing agricultural imaging systems fail when deployed in uncontrolled outdoor lighting conditions.
"""
    p2.insert_text((50, 60), p2_text, fontsize=11)

    # Page 3: Methodology & Future Scope
    p3 = doc.new_page()
    p3_text = """METHODOLOGY
Our proposed system architecture comprises image acquisition, preprocessing with histogram equalization,
and a ResNet-50 classification backbone trained with cross-entropy loss.

RESULTS AND DISCUSSION
The model achieved 91.4% classification accuracy across 10 disease categories on benchmark sets.

FUTURE SCOPE
1. High computational cost prevents real-time deployment on low-cost agricultural drones or Raspberry Pi microcontrollers.
2. Model performance drops significantly when background soil or weeds occlude the leaf margins.
3. Lack of federated edge updates means farmers cannot train local disease varieties without central server access.
"""
    p3.insert_text((50, 60), p3_text, fontsize=11)
    doc.save(str(pdf_path))
    doc.close()

    # 2. CS_1 / Group 1 / Abstract / Synopsis.docx
    grp1_abs_dir = raw_dir / "CS_1" / "Group 1" / "Abstract"
    grp1_abs_dir.mkdir(parents=True, exist_ok=True)
    docx_path = grp1_abs_dir / "Synopsis.docx"
    d = Document()
    d.add_heading("Smart Agriculture Leaf Disease Detection - Synopsis", level=1)
    d.add_paragraph("Submitted by: Vaibhav Kadam (Student ID: PRN72018492K), Amit Deshmukh")
    d.add_heading("ABSTRACT", level=2)
    d.add_paragraph("This synopsis outlines the deep learning framework for leaf anomaly classification.")
    d.save(str(docx_path))

    # 3. CS_1 / Group 1 / Review 4 / Presentation.pptx
    grp1_rev_dir = raw_dir / "CS_1" / "Group 1" / "Review 4"
    grp1_rev_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = grp1_rev_dir / "Review4_Presentation.pptx"
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Smart Agriculture Review 4"
    slide1.placeholders[1].text = "Vaibhav Kadam & Amit Deshmukh\nGuide: Prof. Sharma"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "FUTURE SCOPE & LIMITATIONS"
    slide2.placeholders[1].text = "Edge device latency is high. Model quantization and multi-spectral sensors required for future work."
    prs.save(str(pptx_path))

    # 4. DS_2 / Group 3 / Research Paper / Healthcare_IoT.docx
    grp2_dir = raw_dir / "DS_2" / "Group 3" / "Research Paper"
    grp2_dir.mkdir(parents=True, exist_ok=True)
    doc2_path = grp2_dir / "Healthcare_IoT_Paper.docx"
    d2 = Document()
    d2.add_heading("Healthcare IoT: Real-Time Patient Vital Monitoring", level=0)
    d2.add_paragraph("Submitted by: Sneha Patil, Rohan Joshi. Guide: Dr. R. V. Kulkarni")
    d2.add_heading("ABSTRACT", level=1)
    d2.add_paragraph("Remote patient monitoring wearable capable of streaming ECG and accelerometer telemetry.")
    d2.add_heading("METHODOLOGY", level=1)
    d2.add_paragraph("ESP32 microcontroller with AD8232 ECG sensor module streaming over MQTT.")
    d2.add_heading("FUTURE SCOPE", level=1)
    d2.add_paragraph("Current battery life is limited to 6 hours. Lack of offline edge inference requires further research in TinyML.")
    d2.save(str(doc2_path))


class TestProjectIQEndToEnd(unittest.TestCase):
    def setUp(self):
        self.work_dir = Path(tempfile.mkdtemp(prefix="projectiq_e2e_test_"))
        self.raw_dir = self.work_dir / "raw"
        self.extracted_dir = self.work_dir / "extracted"
        self.cleaned_dir = self.work_dir / "cleaned"
        self.chunks_file = self.work_dir / "chunks" / "chunks.jsonl"
        self.chroma_dir = self.work_dir / "chroma_db"
        self.reports_dir = self.work_dir / "reports"

        # Generate self-contained multi-document corpus
        create_synthetic_college_corpus(self.raw_dir)

    def tearDown(self):
        shutil.rmtree(self.work_dir, ignore_errors=True)

    def test_full_pipeline_run_self_contained(self):
        summary = run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=True,
            domain_filter="Computer Vision & Healthcare Diagnostics"
        )

        # 1. Verify Extraction Stage
        self.assertEqual(summary["extracted_files"], 4)
        manifest_file = self.extracted_dir / "_manifest.json"
        self.assertTrue(manifest_file.exists())
        manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
        self.assertEqual(manifest["extracted_files"], 4)
        self.assertEqual(manifest["failed_files"], 0)

        # 2. Verify Cleaning & Privacy Stage
        self.assertGreaterEqual(summary["cleaned_docs"], 3)
        cleaning_report = self.reports_dir / "cleaning_report.json"
        self.assertTrue(cleaning_report.exists())
        clean_stats = json.loads(cleaning_report.read_text(encoding="utf-8"))
        self.assertIn("total_redactions", clean_stats)

        # 3. Verify Chunking & Token Bounds
        self.assertTrue(self.chunks_file.exists())
        self.assertGreater(summary["total_chunks"], 0)
        with open(self.chunks_file, "r", encoding="utf-8") as f:
            for line in f:
                c = json.loads(line)
                self.assertLessEqual(c["token_count"], 520, f"Token overflow detected: {c['token_count']}")
                self.assertIn("is_canonical", c)
                self.assertIn("category", c)
                self.assertIn("section", c)

        # 4. Verify Indexing in ChromaDB
        self.assertGreater(summary["indexed_chunks"], 0)
        store = ProjectIQVectorStore(db_path=self.chroma_dir)
        self.assertEqual(store.count(), summary["indexed_chunks"])

        # 5. Verify Research Gap Proposals
        self.assertGreater(summary["gaps_found"], 0)
        self.assertGreater(summary["proposals_generated"], 0)
        proposals_file = self.reports_dir / "novel_project_proposals.json"
        self.assertTrue(proposals_file.exists())
        proposals = json.loads(proposals_file.read_text(encoding="utf-8"))
        self.assertEqual(len(proposals), summary["proposals_generated"])
        p0 = proposals[0]
        self.assertIn("title", p0)
        self.assertIn("past_project_citations", p0)
        self.assertIn("feasibility_score", p0)

    def test_pipeline_reset_clears_previous_outputs(self):
        # First run
        run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=False
        )

        # Add a canary file to extracted_dir
        canary = self.extracted_dir / "stale_canary.txt"
        canary.write_text("should be wiped by reset")
        self.assertTrue(canary.exists())

        # Second run with reset=True
        run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=True
        )
        self.assertFalse(canary.exists())

    def test_pipeline_with_iot_agriculture_domain_filter(self):
        summary = run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=True,
            domain_filter="IoT, Edge AI & Smart Agriculture"
        )
        proposals_file = self.reports_dir / "novel_project_proposals.json"
        proposals = json.loads(proposals_file.read_text(encoding="utf-8"))
        self.assertGreater(len(proposals), 0)
        domain_str = proposals[0]["domain"].lower()
        self.assertTrue("agriculture" in domain_str or "iot" in domain_str or "edge" in domain_str)

    def test_pipeline_strict_token_bound_compliance(self):
        run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=True
        )
        with open(self.chunks_file, "r", encoding="utf-8") as f:
            for line in f:
                c = json.loads(line)
                self.assertLessEqual(c["token_count"], 520)

    def test_pipeline_end_to_end_team_recommendation(self):
        run_full_pipeline(
            raw_dir=self.raw_dir,
            extracted_dir=self.extracted_dir,
            cleaned_dir=self.cleaned_dir,
            chunks_file=self.chunks_file,
            chroma_dir=self.chroma_dir,
            reports_dir=self.reports_dir,
            reset=True
        )
        # Evaluate student team with embedded sensors and IoT skills
        matcher = TeamMatcher()
        eval_result = matcher.evaluate_team(["ESP32", "Sensors", "MQTT", "Python"], top_k=1)
        self.assertEqual(len(eval_result["top_domains"]), 1)
        top = eval_result["top_domains"][0]
        self.assertIn("IoT, Edge AI & Smart Agriculture", top["domain"])
        self.assertIn("High", top["team_readiness"])


if __name__ == "__main__":
    unittest.main()
