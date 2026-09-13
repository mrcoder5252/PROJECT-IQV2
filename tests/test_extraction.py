"""
Unit tests for Project-IQ Document Extraction (Tier 1: Feature Coverage)
Covers multi-format parsing (PDF, DOCX, PPTX), SHA-256 hashing, stable IDs,
college hierarchy parsing, text cleaning, and manifest generation.
"""

import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

from src.extract import (
    clean_text,
    extract_all,
    extract_docx,
    extract_one,
    extract_pdf,
    extract_pptx,
    parse_college_hierarchy,
    sha256_file,
    stable_id,
)


class TestExtraction(unittest.TestCase):
    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="projectiq_extract_test_"))

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_parse_college_hierarchy_cs1(self):
        rel = "CS_1/Group 1/Group 1/Blackbook/Acupath_Black_Book_Project.pdf"
        meta = parse_college_hierarchy(rel)
        self.assertEqual(meta["department"], "CS_1")
        self.assertEqual(meta["group_id"], "Group 1")
        self.assertEqual(meta["category"], "blackbook")

    def test_parse_college_hierarchy_ds2_review4(self):
        rel = "DS_2/Group 5/Review 4/Presentation.pptx"
        meta = parse_college_hierarchy(rel)
        self.assertEqual(meta["department"], "DS_2")
        self.assertEqual(meta["group_id"], "Group 5")
        self.assertEqual(meta["category"], "review4")

    def test_clean_text(self):
        raw = "Hello\x00 world!\n\n\n\nTest   spacing."
        cleaned = clean_text(raw)
        self.assertEqual(cleaned, "Hello world!\n\nTest spacing.")

    def test_clean_text_none_and_soft_hyphens(self):
        self.assertEqual(clean_text(None), "")
        self.assertEqual(clean_text("co\u00adordi\u00adnator"), "coordinator")

    def test_stable_id_deterministic(self):
        id1 = stable_id("doc1", 0, "paragraph")
        id2 = stable_id("doc1", 0, "paragraph")
        id3 = stable_id("doc1", 1, "paragraph")
        self.assertEqual(id1, id2)
        self.assertNotEqual(id1, id3)
        self.assertEqual(len(id1), 40)  # SHA-1 40-character hex string

    def test_sha256_file(self):
        sample_file = self.test_dir / "sample.bin"
        test_bytes = b"Project-IQ academic document binary content test 12345"
        sample_file.write_bytes(test_bytes)

        expected_hash = hashlib.sha256(test_bytes).hexdigest()
        actual_hash = sha256_file(sample_file)
        self.assertEqual(actual_hash, expected_hash)
        self.assertEqual(len(actual_hash), 64)

    def test_extract_pdf_synthetic(self):
        pdf_path = self.test_dir / "synthetic_report.pdf"
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((50, 72), "CHAPTER 1 INTRODUCTION\nThis is a synthetic academic report.")
        p2 = doc.new_page()
        p2.insert_text((50, 72), "METHODOLOGY\nDeep learning models were trained on GPU.")
        doc.save(str(pdf_path))
        doc.close()

        blocks = extract_pdf(pdf_path)
        self.assertGreaterEqual(len(blocks), 2)
        texts = [b["text"] for b in blocks]
        self.assertTrue(any("INTRODUCTION" in t for t in texts))
        self.assertTrue(any("METHODOLOGY" in t for t in texts))
        self.assertIn("page", blocks[0])
        self.assertIn("block_type", blocks[0])

    def test_extract_docx_synthetic(self):
        docx_path = self.test_dir / "synthetic_paper.docx"
        doc = Document()
        doc.add_heading("Deep Learning for Leaf Disease Detection", level=1)
        doc.add_paragraph("Abstract: Plant leaf disease causes major agricultural losses.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Model"
        table.cell(0, 1).text = "Accuracy"
        table.cell(1, 0).text = "ResNet-50"
        table.cell(1, 1).text = "94.5%"
        doc.save(str(docx_path))

        blocks = extract_docx(docx_path)
        self.assertGreaterEqual(len(blocks), 2)
        block_types = [b["block_type"] for b in blocks]
        self.assertIn("heading", block_types)
        self.assertIn("paragraph", block_types)
        self.assertIn("table", block_types)

    def test_extract_pptx_synthetic(self):
        pptx_path = self.test_dir / "synthetic_presentation.pptx"
        prs = Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Review 4 Presentation"
        slide1.placeholders[1].text = "Decentralized Blockchain Registry"
        prs.save(str(pptx_path))

        blocks = extract_pptx(pptx_path)
        self.assertGreaterEqual(len(blocks), 1)
        texts = [b["text"] for b in blocks]
        self.assertTrue(any("Review 4" in t for t in texts))

    def test_extract_all_and_manifest(self):
        raw_dir = self.test_dir / "raw" / "CS_1" / "Group 1" / "Blackbook"
        raw_dir.mkdir(parents=True, exist_ok=True)
        out_dir = self.test_dir / "extracted"

        # Create 1 synthetic PDF
        pdf_path = raw_dir / "Project_Report.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 72), "FINAL YEAR PROJECT REPORT\nAcupath Pathology Diagnostics.")
        doc.save(str(pdf_path))
        doc.close()

        # Create 1 synthetic DOCX
        docx_path = raw_dir / "Synopsis.docx"
        d = Document()
        d.add_paragraph("Synopsis text for group 1 project.")
        d.save(str(docx_path))

        manifest = extract_all(input_dir=self.test_dir / "raw", output_dir=out_dir)
        self.assertEqual(manifest["extracted_files"], 2)
        self.assertEqual(manifest["failed_files"], 0)
        self.assertEqual(len(manifest["documents"]), 2)

        manifest_file = out_dir / "_manifest.json"
        self.assertTrue(manifest_file.exists())
        saved_manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
        self.assertEqual(saved_manifest["extracted_files"], 2)
        self.assertIn("documents", saved_manifest)


if __name__ == "__main__":
    unittest.main()
