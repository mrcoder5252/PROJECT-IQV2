"""
Project-IQ: Multi-Format Academic Document Extractor
Supports PDF, PPTX, DOCX, and DOC documents with OCR fallback.
Understands college project hierarchy (Department / Group / Document Type).
Generates structured block-level JSON outputs with persistent SHA-256 hashes and stable IDs.
"""

from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Dict, List, Tuple

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from PIL import Image

try:
    import pytesseract
    _HAS_TESSERACT = True
except ImportError:
    _HAS_TESSERACT = False

from src.config import config
from src.logger import get_logger

logger = get_logger("project_iq.extract")

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".doc"}


def sha256_file(path: Path) -> str:
    """Compute streamed SHA-256 binary hash of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_id(*parts: Any) -> str:
    """Generate deterministic stable block ID using SHA-1 hex digest of serialized parts."""
    return hashlib.sha1("|".join(map(str, parts)).encode("utf-8")).hexdigest()


def clean_text(s: str | None) -> str:
    """Clean and normalize raw extracted text strings."""
    if not s:
        return ""
    s = s.replace("\x00", " ").replace("\u00ad", "")
    s = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", s)
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def parse_college_hierarchy(rel_path: str, full_path: Path | None = None) -> Dict[str, str]:
    """Parse department, group ID, category, and project name from relative/full path."""
    path_str = str(full_path or rel_path).replace("\\", "/")
    parts = [p for p in rel_path.replace("\\", "/").split("/") if p]
    meta = {
        "department": "General",
        "group_id": "General",
        "category": "project_report",
        "project_name": "Unknown"
    }

    # Department detection from full path or rel path
    dept_match = re.search(r"\b(CS_\d+|DS_\d+|IT_\d+|AI_\d+|COMP_\d+|EXTC_\d+)\b", path_str, re.I)
    if dept_match:
        meta["department"] = dept_match.group(1).upper()

    for p in parts:
        if re.match(r"^Group\s*\d+$", p, re.I):
            meta["group_id"] = p.title()
            break

    # Detect category from folder hierarchy or filename
    lower_path = rel_path.lower()
    if any(k in lower_path for k in ["blackbook", "black_book", "black book", "final_report", "project_report"]):
        meta["category"] = "blackbook"
    elif any(k in lower_path for k in ["research_paper", "paper", "ieee", "springer", "journal", "conference"]):
        meta["category"] = "research_paper"
    elif any(k in lower_path for k in ["abstract", "synopsis"]):
        meta["category"] = "abstract_synopsis"
    elif any(k in lower_path for k in ["review 4", "review4", "review_4", "review no 4", "presentation", "ppt", "poster"]):
        meta["category"] = "review4"

    filename_stem = Path(parts[-1]).stem
    clean_stem = re.sub(
        r"(?i)(final|report|project|blackbook|synopsis|abstract|review|100%|poster|presentation|\(|\)|\d+)",
        "",
        filename_stem
    ).strip(" _-")
    meta["project_name"] = clean_stem.title() if len(clean_stem) >= 3 else filename_stem

    return meta


def extract_front_matter_metadata(
    blocks: List[Dict[str, Any]],
    hierarchy: Dict[str, str],
    path: Path
) -> Dict[str, Any]:
    """
    Extract structured metadata (title, authors, year, supervisor, category)
    from the front-matter blocks of an academic document.
    """
    front_blocks = blocks[:15]
    front_text = "\n".join(b.get("text", "") for b in front_blocks)

    # 1. Title Extraction
    title = ""
    title_match = re.search(
        r"(?i)(?:A\s+PROJECT\s+REPORT\s+ON|A\s+SEMINAR\s+REPORT\s+ON|PROJECT\s+TITLE\s*[:\-]|TITLE\s*[:\-])\s*\n*([^\n\r]+(?:\n[^\n\r]+)?)",
        front_text
    )
    if title_match:
        cand = clean_text(title_match.group(1))
        cand = re.split(r"(?i)\b(?:submitted\s+by|guided\s+by|by\s*:)\b", cand)[0].strip()
        if len(cand) >= 5 and not any(k in cand.lower() for k in ["university", "department", "college"]):
            title = cand

    # If DOCX: check heading style blocks
    if not title:
        for b in front_blocks:
            if b.get("block_type") == "heading":
                cand = clean_text(b.get("text", ""))
                if cand and len(cand) >= 5 and not any(k in cand.lower() for k in ["abstract", "acknowledgement", "certificate", "contents", "table of"]):
                    title = cand
                    break

    # If PPTX: first slide title or first block on slide 1
    if not title:
        for b in front_blocks:
            if b.get("slide") == 1 and b.get("text"):
                cand = clean_text(b.get("text", ""))
                first_line = cand.split("\n")[0].strip()
                if len(first_line) >= 5 and not any(k in first_line.lower() for k in ["review", "presentation", "welcome"]):
                    title = first_line
                    break

    # Search page 1 blocks for candidate title lines
    if not title:
        for b in front_blocks:
            if b.get("page", 1) == 1:
                for line in b.get("text", "").splitlines():
                    cl = clean_text(line)
                    if 15 <= len(cl) <= 160:
                        low = cl.lower()
                        if not any(k in low for k in [
                            "university", "department", "submitted by", "guided by", "pune", "mumbai",
                            "college", "roll no", "prn", "certificate", "report on", "phone", "email"
                        ]):
                            title = cl
                            break
            if title:
                break

    # Fallback to hierarchy project name or cleaned file stem
    if not title:
        clean_name = hierarchy.get("project_name", "")
        if clean_name and clean_name != "Unknown":
            title = clean_name
        else:
            title = path.stem.replace("_", " ").title()

    # 2. Authors Extraction
    authors: List[str] = []
    author_match = re.search(
        r"(?i)(?:SUBMITTED\s+BY|TEAM\s+MEMBERS?|AUTHORS?|STUDENTS?|PREPARED\s+BY)\s*[:\-]?\s*([\s\S]+?)(?=(?:GUIDED\s+BY|GUIDE|UNDER\s+THE\s+GUIDANCE|SUPERVISOR|DEPARTMENT|COLLEGE|THIS\s+IS\s+TO\s+CERTIFY|ABSTRACT|ACKNOWLEDGEMENT|\n\s*\n\s*\n|$))",
        front_text
    )
    if author_match:
        raw_author_block = author_match.group(1)
        for line in raw_author_block.splitlines():
            line_str = clean_text(line)
            if not line_str or any(k in line_str.lower() for k in ["email", "@", "phone", "mobile", "contact"]):
                continue
            cleaned_line = re.sub(r"\([^)]*\)", "", line_str)
            cleaned_line = re.sub(r"(?i)\b(?:roll\s*no|prn|student\s*id|seat\s*no)\s*[:#=\-]?\s*[a-z0-9\-_/]+", "", cleaned_line)
            parts = re.split(r"[,&]|\band\b", cleaned_line)
            for p in parts:
                name = clean_text(p).strip(" -:;")
                name_words = name.split()
                if 2 <= len(name_words) <= 4 and 4 <= len(name) <= 40:
                    if not any(k in name.lower() for k in [
                        "department", "university", "college", "guided", "guide",
                        "supervisor", "submitted", "roll", "prn"
                    ]):
                        if name not in authors:
                            authors.append(name)

    # 3. Supervisor Extraction
    supervisor = None
    sup_match = re.search(
        r"(?i)(?:GUIDED\s+BY|GUIDE|SUPERVISOR|PROJECT\s+GUIDE|PROJECT\s+MENTOR|FACULTY\s+GUIDE|UNDER\s+THE\s+GUIDANCE\s+OF)\s*[:\-]?\s*([^\n\r,]+)",
        front_text
    )
    if sup_match:
        cand = clean_text(sup_match.group(1)).strip(" -:;")
        if 3 <= len(cand) <= 50 and not any(k in cand.lower() for k in ["department", "university", "college", "submitted"]):
            supervisor = cand

    # 4. Year Extraction
    year = None
    year_match = re.search(r"\b(20\d{2}[\s\-_/]+(?:20)?\d{2})\b", front_text)
    if year_match:
        year = year_match.group(1).replace(" ", "")
    else:
        single_year_match = re.search(r"\b(20\d{2})\b", front_text)
        if single_year_match:
            year = single_year_match.group(1)

    # 5. Category
    category = hierarchy.get("category", "project_report")

    return {
        "title": title,
        "authors": authors,
        "year": year,
        "supervisor": supervisor,
        "category": category
    }


def extract_pdf(path: Path) -> List[Dict[str, Any]]:
    """
    Extract text blocks from PDF using PyMuPDF (fitz) with PyTesseract OCR fallback
    for scanned or image-heavy pages (< 30 characters).
    """
    out = []
    doc = fitz.open(path)
    try:
        for page_no, page in enumerate(doc, 1):
            page_blocks = []
            for block_no, b in enumerate(page.get_text("blocks")):
                text = clean_text(b[4])
                if text:
                    page_blocks.append({
                        "block_type": "paragraph",
                        "text": text,
                        "page": page_no,
                        "block_no": block_no,
                        "char_count": len(text),
                        "ocr": False
                    })
            chars = sum(len(x["text"]) for x in page_blocks)
            out.extend(page_blocks)

            # OCR fallback for scanned pages or image-only PDFs (< 30 chars)
            if chars < 30 and _HAS_TESSERACT:
                try:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr = clean_text(pytesseract.image_to_string(img))
                    if ocr:
                        out.append({
                            "block_type": "ocr",
                            "text": ocr,
                            "page": page_no,
                            "block_no": 0,
                            "char_count": len(ocr),
                            "ocr": True,
                            "ocr_quality": "fallback"
                        })
                except Exception as e:
                    logger.debug(f"OCR fallback skipped on page {page_no}: {e}")
    finally:
        doc.close()
    return out


def extract_pptx(path: Path) -> List[Dict[str, Any]]:
    """Extract text and tables from PPTX slides."""
    out = []
    prs = Presentation(path)
    for slide_no, slide in enumerate(prs.slides, 1):
        found = False
        for shape_no, shape in enumerate(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                text = clean_text(shape.text)
                if text:
                    found = True
                    out.append({
                        "block_type": "paragraph",
                        "text": text,
                        "slide": slide_no,
                        "page": slide_no,
                        "shape_no": shape_no,
                        "char_count": len(text),
                        "ocr": False
                    })
            if getattr(shape, "has_table", False):
                rows = [[clean_text(c.text) for c in row.cells] for row in shape.table.rows]
                if any(any(c for c in row) for row in rows):
                    found = True
                    table_text = "\n".join(" | ".join(row) for row in rows)
                    out.append({
                        "block_type": "table",
                        "text": table_text,
                        "rows": rows,
                        "slide": slide_no,
                        "page": slide_no,
                        "shape_no": shape_no,
                        "char_count": len(table_text),
                        "ocr": False
                    })
        if not found:
            out.append({
                "block_type": "slide_placeholder",
                "text": "",
                "slide": slide_no,
                "page": slide_no,
                "shape_no": -1,
                "char_count": 0,
                "ocr": False,
                "image_only": True
            })
    return out


def extract_docx(path: Path) -> List[Dict[str, Any]]:
    """Extract paragraphs and tables from DOCX documents."""
    out = []
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs):
        text = clean_text(p.text)
        if text:
            is_heading = "heading" in p.style.name.lower()
            out.append({
                "block_type": "heading" if is_heading else "paragraph",
                "text": text,
                "paragraph_no": i,
                "page": 1,
                "char_count": len(text),
                "ocr": False
            })
    for table_no, table in enumerate(doc.tables):
        rows = [[clean_text(c.text) for c in row.cells] for row in table.rows]
        if any(any(c for c in row) for row in rows):
            t_text = "\n".join(" | ".join(row) for row in rows)
            out.append({
                "block_type": "table",
                "text": t_text,
                "rows": rows,
                "table_no": table_no,
                "page": 1,
                "char_count": len(t_text),
                "ocr": False
            })
    return out


def convert_doc(path: Path) -> Tuple[Path, Path]:
    """Convert legacy .doc document to .docx using headless LibreOffice."""
    tmp = Path(tempfile.mkdtemp(prefix="projectiq_doc_"))
    try:
        lo_bin = shutil.which("libreoffice") or shutil.which("soffice")
        if not lo_bin:
            # Check standard Windows installation directories
            win_candidates = [
                Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")
            ]
            for candidate in win_candidates:
                if candidate.exists():
                    lo_bin = str(candidate)
                    break
        if not lo_bin:
            raise RuntimeError("LibreOffice binary (libreoffice/soffice) not found on system PATH.")

        subprocess.run(
            [lo_bin, "--headless", "--convert-to", "docx",
             "--outdir", str(tmp), str(path)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60
        )
        converted = tmp / f"{path.stem}.docx"
        if not converted.exists():
            raise RuntimeError(f"Could not convert {path} to DOCX")
        return converted, tmp
    except Exception as e:
        shutil.rmtree(tmp, ignore_errors=True)
        raise RuntimeError(f"LibreOffice conversion failed for {path}: {e}")


def extract_one(path: Path, root: Path) -> Dict[str, Any]:
    """Extract a single document, compute SHA-256 and stable doc/block IDs."""
    tmp_dir = None
    actual_path = path
    ext = path.suffix.lower()

    try:
        if ext == ".doc":
            actual_path, tmp_dir = convert_doc(path)
            blocks = extract_docx(actual_path)
            dtype = "doc"
        elif ext == ".pdf":
            blocks = extract_pdf(actual_path)
            dtype = "pdf"
        elif ext == ".pptx":
            blocks = extract_pptx(actual_path)
            dtype = "pptx"
        elif ext == ".docx":
            blocks = extract_docx(actual_path)
            dtype = "docx"
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

        rel = str(path.relative_to(root)).replace("\\", "/")
        sha = sha256_file(path)
        doc_id = f"doc_{sha[:12]}"
        hierarchy = parse_college_hierarchy(rel, path)
        metadata = extract_front_matter_metadata(blocks, hierarchy, path)
        now_utc = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

        result_blocks = []
        for i, b in enumerate(blocks):
            b_id = stable_id(doc_id, i, b.get("block_type", "paragraph"))
            result_blocks.append({
                "block_id": b_id,
                "document_id": doc_id,
                "doc_id": doc_id,
                "document_version": sha[:16],
                "file_sha256": sha,
                "sha256": sha,
                "source_file": path.name,
                "source_path": rel,
                "source_index": i,
                "department": hierarchy["department"],
                "group_id": hierarchy["group_id"],
                "category": hierarchy["category"],
                "document_type": dtype,
                **b
            })

        doc_meta = {
            "doc_id": doc_id,
            "document_id": doc_id,
            "document_version": sha[:16],
            "file_sha256": sha,
            "sha256": sha,
            "source_file": path.name,
            "source_path": rel,
            "department": hierarchy["department"],
            "group_id": hierarchy["group_id"],
            "project_name": hierarchy["project_name"],
            "category": hierarchy["category"],
            "document_type": dtype,
            "extracted_at": now_utc,
            "metadata": metadata,
            "block_count": len(result_blocks)
        }

        return {
            "doc_id": doc_id,
            "document_id": doc_id,
            "source_file": path.name,
            "source_path": rel,
            "sha256": sha,
            "file_sha256": sha,
            "category": hierarchy["category"],
            "extracted_at": now_utc,
            "metadata": metadata,
            "document": doc_meta,
            "blocks": result_blocks
        }
    finally:
        if tmp_dir and Path(tmp_dir).exists():
            shutil.rmtree(tmp_dir, ignore_errors=True)


def extract_all(input_dir: str | Path, output_dir: str | Path) -> Dict[str, Any]:
    """
    Extract all supported documents from input_dir into output_dir.
    Writes individual {doc_id}.json records and a cataloging _manifest.json.
    """
    root = Path(input_dir).resolve()
    out = Path(output_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)

    manifest: List[Dict[str, Any]] = []
    seen_hashes: Dict[str, str] = {}
    successful = 0
    failed = 0

    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            data = extract_one(path, root)
            meta = data["document"]
            sha = meta["file_sha256"]

            if sha in seen_hashes:
                meta["is_duplicate_file"] = True
                meta["duplicate_of"] = seen_hashes[sha]
            else:
                meta["is_duplicate_file"] = False
                meta["duplicate_of"] = None
                seen_hashes[sha] = meta["document_id"]

            out_file = out / f"{meta['document_id']}.json"
            out_file.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            manifest.append(meta)
            successful += 1
            logger.info(f"EXTRACTED [{meta['department']}/{meta['group_id']} | {meta['category']}]: {path.name}")
        except Exception as e:
            failed += 1
            logger.error(f"FAILED: {path.name} -> {repr(e)}")

    manifest_data = {
        "total_files": successful + failed,
        "extracted_files": successful,
        "failed_files": failed,
        "unique_content_files": len(seen_hashes),
        "documents": manifest
    }
    (out / "_manifest.json").write_text(
        json.dumps(manifest_data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest_data


def main():
    parser = argparse.ArgumentParser(description="Extract text and structure from academic project documents.")
    parser.add_argument("--input", required=True, help="Input directory containing raw documents")
    parser.add_argument("--output", required=True, help="Output directory for extracted JSON records")
    args = parser.parse_args()

    manifest = extract_all(args.input, args.output)
    logger.info(f"\nEXTRACTION COMPLETE: {manifest['extracted_files']} files extracted, {manifest['failed_files']} failed.")


if __name__ == "__main__":
    main()
